#!/usr/bin/env python3
"""
Parse Excel file and generate JSON data for the quiz application
"""
import openpyxl
import json

def is_yellow_cell(cell):
    """Check if a cell has yellow background"""
    if cell.fill and cell.fill.start_color:
        color = cell.fill.start_color.rgb
        if color and color in ['FFFFFF00', 'FFFF00', '00FFFF00']:
            return True
    return False

def parse_excel_to_json(excel_file):
    """Parse Excel file and extract questions with correct answers"""
    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    
    questions = []
    
    # Skip header row, start from row 2
    for row_idx in range(2, ws.max_row + 1):
        question_id = ws.cell(row=row_idx, column=1).value
        topic = ws.cell(row=row_idx, column=2).value
        question_text = ws.cell(row=row_idx, column=3).value
        
        # Get answers (columns 4-9: A, B, C, D, E, F)
        answers = []
        correct_answers = []
        
        for col_idx in range(4, 10):  # Columns D through I (A through F answers)
            cell = ws.cell(row=row_idx, column=col_idx)
            answer_text = cell.value
            
            if answer_text:  # Only add non-empty answers
                answer_label = chr(65 + col_idx - 4)  # A, B, C, D, E, F
                answers.append({
                    'label': answer_label,
                    'text': str(answer_text)
                })
                
                if is_yellow_cell(cell):
                    correct_answers.append(answer_label)
        
        if question_text and answers:  # Only add valid questions
            questions.append({
                'id': question_id,
                'topic': topic,
                'question': question_text,
                'answers': answers,
                'correct': correct_answers
            })
    
    return questions

if __name__ == '__main__':
    questions = parse_excel_to_json('Database_Tests_ALL_QUESTIONS.xlsx')
    
    # Save to JSON file
    with open('questions.json', 'w', encoding='utf-8') as f:
        json.dump(questions, f, indent=2, ensure_ascii=False)
    
    print(f"Parsed {len(questions)} questions successfully!")
    print(f"Questions saved to questions.json")
